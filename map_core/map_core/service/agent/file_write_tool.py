from __future__ import annotations

import asyncio
import csv
import json
import mimetypes
import os
from datetime import UTC, datetime
from pathlib import Path
from typing import Any, Literal
from uuid import uuid4

from pydantic import BaseModel, ConfigDict, Field, ValidationError

from ...schema.attachment_schema import AttachmentSchema
from ..attachment_collector import AttachmentCollector
from .base import AgentRequest
from .tool_call_agent import Tool

TOOL_NAME = "attachment_file_write_tool"
_DEFAULT_TIMEOUT_SECONDS = 20.0
_DEFAULT_OUTPUT_DIR = "/tmp/map_generated_files"
_MINIO_BASE_URL = os.getenv("MAP_MINIO_BASE_URL", "http://10.16.11.40:9000")
_MINIO_ENDPOINT = (
    _MINIO_BASE_URL.removeprefix("http://").removeprefix("https://").rstrip("/")
)
_MINIO_ACCESS_KEY = os.getenv("MAP_MINIO_ACCESS_KEY", "minioadmin")
_MINIO_SECRET_KEY = os.getenv("MAP_MINIO_SECRET_KEY", "minioadmin")
_MINIO_BUCKET = os.getenv("MAP_MINIO_BUCKET", "map_core")
_MINIO_SECURE_ENV = os.getenv("MAP_MINIO_SECURE")
_MINIO_SECURE = (
    _MINIO_SECURE_ENV.strip().lower() in {"1", "true", "yes", "on"}
    if _MINIO_SECURE_ENV
    else _MINIO_BASE_URL.startswith("https://")
)
_MINIO_PUBLIC_READ = True
_MAX_FILES_PER_REQUEST = 20

_TEXT_FILE_TYPES = {
    "txt",
    "text",
    "md",
    "markdown",
    "json",
    "jsonm",
    "jsonl",
    "log",
    "yaml",
    "yml",
    "xml",
}
_CSV_FILE_TYPES = {"csv"}
_EXCEL_FILE_TYPES = {"xlsx", "excel"}
_WORD_FILE_TYPES = {"docx", "word"}


class FileWriteItemInput(BaseModel):
    model_config = ConfigDict(extra="forbid")

    file_name: str = Field(..., min_length=1, max_length=255)
    file_type: Literal[
        "txt",
        "text",
        "md",
        "markdown",
        "json",
        "jsonm",
        "jsonl",
        "log",
        "yaml",
        "yml",
        "xml",
        "csv",
        "xlsx",
        "excel",
        "docx",
        "word",
    ]
    content: Any
    file_id: str | None = None
    csv_headers: list[str] | None = None
    sheet_name: str = Field(default="Sheet1", min_length=1, max_length=31)


class FileWriteToolInput(BaseModel):
    model_config = ConfigDict(extra="forbid")

    files: list[FileWriteItemInput] = Field(
        ..., min_length=1, max_length=_MAX_FILES_PER_REQUEST
    )
    timeout_seconds: float = Field(default=_DEFAULT_TIMEOUT_SECONDS, gt=0, le=120)
    output_dir: str = Field(default=_DEFAULT_OUTPUT_DIR, min_length=1)
    minio_object_prefix: str = Field(default="generated", min_length=1)
    upload_to_minio: bool = True


class SavedFileOutput(BaseModel):
    model_config = ConfigDict(extra="forbid")

    success: bool
    file_id: str
    file_name: str
    file_type: str
    file_url: str
    local_path: str
    minio_object_key: str
    header: str
    error: str | None = None


class FileWriteStatsOutput(BaseModel):
    model_config = ConfigDict(extra="forbid")

    total_files: int
    success_files: int
    failed_files: int


class FileWriteToolOutput(BaseModel):
    model_config = ConfigDict(extra="forbid")

    success: bool
    message: str
    combined_text: str
    attachments: list[AttachmentSchema]
    files: list[SavedFileOutput]
    stats: FileWriteStatsOutput
    error: str | None = None


def _build_attachment_message(attachments: list[AttachmentSchema]) -> str:
    if not attachments:
        return "未生成任何文件。"
    file_names = ", ".join(item.file_name for item in attachments)
    return f"已生成 {len(attachments)} 个文件：{file_names}"


def _extract_attachments_from_saved_files(
    files: list[SavedFileOutput],
) -> list[AttachmentSchema]:
    attachments: list[AttachmentSchema] = []
    for item in files:
        if not item.success:
            continue
        attachments.append(
            AttachmentSchema(
                file_id=item.file_id,
                file_name=item.file_name,
                file_type=item.file_type,
                file_url=item.file_url,
            )
        )
    return attachments


def _normalize_file_name(file_name: str, file_type: str) -> str:
    safe_name = Path(file_name).name.strip() or "attachment"
    ext = Path(safe_name).suffix.lower().lstrip(".")
    wanted = {
        "text": "txt",
        "markdown": "md",
        "excel": "xlsx",
        "word": "docx",
    }.get(file_type, file_type)
    if ext != wanted:
        safe_name = f"{Path(safe_name).stem}.{wanted}"
    return safe_name


def _build_minio_url(bucket_name: str, object_key: str) -> str:
    if not _MINIO_BASE_URL and not _MINIO_ENDPOINT:
        return ""

    base_url = _MINIO_BASE_URL.rstrip("/")
    if not base_url:
        protocol = "https" if _MINIO_SECURE else "http"
        base_url = f"{protocol}://{_MINIO_ENDPOINT}"
    return f"{base_url}/{bucket_name.strip('/')}/{object_key.lstrip('/')}"


def _format_file_header(
    file_id: str, file_name: str, file_type: str, file_url: str
) -> str:
    return (
        f"file_id={file_id or '-'} | "
        f"file_name={file_name or '-'} | "
        f"file_type={file_type or '-'} | "
        f"file_url={file_url or '-'}"
    )


def _serialize_text_content(file_type: str, content: Any) -> str:
    if isinstance(content, str):
        return content
    if file_type in {"json", "jsonm", "jsonl"}:
        return json.dumps(content, ensure_ascii=False, indent=2)
    if isinstance(content, (dict, list)):
        return json.dumps(content, ensure_ascii=False, indent=2)
    return str(content)


def _write_text_file(path: Path, file_type: str, content: Any) -> None:
    text = _serialize_text_content(file_type, content)
    path.write_text(text, encoding="utf-8")


def _rows_from_content_for_table(content: Any) -> list[list[Any]]:
    if isinstance(content, list) and all(
        isinstance(row, (list, tuple)) for row in content
    ):
        return [list(row) for row in content]
    raise ValueError("Table content must be list[dict] or list[list].")


def _dict_rows_and_headers(
    content: Any, csv_headers: list[str] | None = None
) -> tuple[list[dict[str, Any]], list[str]]:
    if not (
        isinstance(content, list) and all(isinstance(row, dict) for row in content)
    ):
        raise ValueError("Dict table content must be list[dict].")

    dict_rows = [dict(row) for row in content]
    headers: list[str] = list(csv_headers or [])
    if not headers:
        for row in dict_rows:
            for key in row:
                if key not in headers:
                    headers.append(str(key))
    return dict_rows, headers


def _write_csv_file(path: Path, content: Any, csv_headers: list[str] | None) -> None:
    if isinstance(content, str):
        path.write_text(content, encoding="utf-8")
        return

    if isinstance(content, list) and all(isinstance(row, dict) for row in content):
        rows, headers = _dict_rows_and_headers(content, csv_headers)
        with path.open("w", encoding="utf-8", newline="") as fp:
            writer = csv.DictWriter(fp, fieldnames=headers)
            writer.writeheader()
            for row in rows:
                writer.writerow({key: row.get(key, "") for key in headers})
        return

    rows = _rows_from_content_for_table(content)
    with path.open("w", encoding="utf-8", newline="") as fp:
        writer = csv.writer(fp)
        writer.writerows(rows)


def _write_excel_file(path: Path, content: Any, sheet_name: str) -> None:
    try:
        from openpyxl import Workbook
    except ImportError as exc:  # pragma: no cover - depends on runtime environment
        raise RuntimeError("`openpyxl` is required to write xlsx/excel files.") from exc

    wb = Workbook()
    ws = wb.active
    if (
        ws is None
    ):  # pragma: no cover - defensive guard for type checkers/runtime safety
        raise RuntimeError("Failed to initialize worksheet.")
    ws.title = sheet_name

    if isinstance(content, list) and all(isinstance(row, dict) for row in content):
        rows, headers = _dict_rows_and_headers(content)
        ws.append(headers)
        for row in rows:
            ws.append([row.get(key, "") for key in headers])
    else:
        rows = _rows_from_content_for_table(content)
        for row in rows:
            ws.append(row)

    wb.save(path)


def _write_word_file(path: Path, content: Any) -> None:
    try:
        from docx import Document
    except ImportError as exc:
        raise RuntimeError(
            "`python-docx` is required to write docx/word files."
        ) from exc

    doc = Document()
    if isinstance(content, str):
        lines = content.splitlines() or [content]
    elif isinstance(content, list):
        lines = [str(item) for item in content]
    else:
        lines = [json.dumps(content, ensure_ascii=False, indent=2)]

    for line in lines:
        doc.add_paragraph(line)
    doc.save(str(path))


def _write_single_file(item: FileWriteItemInput, output_dir: Path) -> tuple[str, Path]:
    file_id = item.file_id or uuid4().hex
    file_type = item.file_type
    file_name = _normalize_file_name(item.file_name, file_type)
    target_path = output_dir / f"{file_id}_{file_name}"

    if file_type in _TEXT_FILE_TYPES:
        _write_text_file(target_path, file_type, item.content)
    elif file_type in _CSV_FILE_TYPES:
        _write_csv_file(target_path, item.content, item.csv_headers)
    elif file_type in _EXCEL_FILE_TYPES:
        _write_excel_file(target_path, item.content, item.sheet_name)
    elif file_type in _WORD_FILE_TYPES:
        _write_word_file(target_path, item.content)
    else:  # pragma: no cover - guarded by pydantic literal
        raise ValueError(f"Unsupported file_type: {file_type}")

    return file_id, target_path


def _build_object_key(prefix: str, file_name: str) -> str:
    date_path = datetime.now(UTC).strftime("%Y/%m/%d")
    normalized_prefix = prefix.strip("/")
    return f"{normalized_prefix}/{date_path}/{file_name}"


def _build_public_read_policy(bucket_name: str) -> str:
    return json.dumps(
        {
            "Version": "2012-10-17",
            "Statement": [
                {
                    "Effect": "Allow",
                    "Principal": {"AWS": ["*"]},
                    "Action": ["s3:GetObject"],
                    "Resource": [f"arn:aws:s3:::{bucket_name}/*"],
                }
            ],
        }
    )


def _upload_to_minio(local_path: Path, object_key: str, upload_to_minio: bool) -> str:
    if not upload_to_minio:
        return ""

    if not _MINIO_ENDPOINT:
        raise RuntimeError("MinIO endpoint is empty. Check MAP_MINIO_BASE_URL.")

    try:
        from minio import Minio
    except ImportError as exc:
        raise RuntimeError("`minio` is required to upload files to MinIO.") from exc

    client = Minio(
        _MINIO_ENDPOINT,
        access_key=_MINIO_ACCESS_KEY,
        secret_key=_MINIO_SECRET_KEY,
        secure=_MINIO_SECURE,
    )

    if not client.bucket_exists(_MINIO_BUCKET):
        client.make_bucket(_MINIO_BUCKET)

    if _MINIO_PUBLIC_READ:
        client.set_bucket_policy(
            _MINIO_BUCKET, _build_public_read_policy(_MINIO_BUCKET)
        )

    content_type = (
        mimetypes.guess_type(local_path.name)[0] or "application/octet-stream"
    )
    client.fput_object(
        _MINIO_BUCKET,
        object_key,
        str(local_path),
        content_type=content_type,
    )
    return _build_minio_url(_MINIO_BUCKET, object_key)


def create_attachment_file_write_tool() -> Tool:
    async def _handler(
        args: dict[str, Any], _request: Any, _parid: str
    ) -> dict[str, Any]:
        try:
            payload = FileWriteToolInput.model_validate(args)
        except ValidationError as exc:
            return {"success": False, "error": exc.errors()}

        def _write_sync() -> FileWriteToolOutput:
            output_dir = Path(payload.output_dir).expanduser()
            output_dir.mkdir(parents=True, exist_ok=True)

            results: list[SavedFileOutput] = []
            combined_parts: list[str] = []

            for item in payload.files:
                try:
                    file_id, local_path = _write_single_file(item, output_dir)
                    object_key = _build_object_key(
                        payload.minio_object_prefix, local_path.name
                    )
                    file_url = _upload_to_minio(
                        local_path, object_key, payload.upload_to_minio
                    )
                    if not file_url:
                        file_url = str(local_path)
                    header = _format_file_header(
                        file_id=file_id,
                        file_name=local_path.name,
                        file_type=item.file_type,
                        file_url=file_url,
                    )
                    combined_parts.append(f"[SavedAttachment]\n{header}")
                    results.append(
                        SavedFileOutput(
                            success=True,
                            file_id=file_id,
                            file_name=local_path.name,
                            file_type=item.file_type,
                            file_url=file_url,
                            local_path=str(local_path),
                            minio_object_key=object_key,
                            header=header,
                        )
                    )
                except Exception as exc:
                    fallback_id = item.file_id or ""
                    fallback_name = _normalize_file_name(item.file_name, item.file_type)
                    header = _format_file_header(
                        file_id=fallback_id,
                        file_name=fallback_name,
                        file_type=item.file_type,
                        file_url="",
                    )
                    results.append(
                        SavedFileOutput(
                            success=False,
                            file_id=fallback_id,
                            file_name=fallback_name,
                            file_type=item.file_type,
                            file_url="",
                            local_path="",
                            minio_object_key="",
                            header=header,
                            error=str(exc),
                        )
                    )

            success_files = sum(1 for item in results if item.success)
            attachments = _extract_attachments_from_saved_files(results)
            stats = FileWriteStatsOutput(
                total_files=len(results),
                success_files=success_files,
                failed_files=len(results) - success_files,
            )
            return FileWriteToolOutput(
                success=success_files > 0,
                message=_build_attachment_message(attachments),
                combined_text="\n\n".join(combined_parts),
                attachments=attachments,
                files=results,
                stats=stats,
            )

        try:
            result = await asyncio.wait_for(
                asyncio.to_thread(_write_sync), timeout=payload.timeout_seconds
            )
            request_extra = getattr(_request, "extra", None)
            if isinstance(request_extra, dict):
                collector = request_extra.get("attachment_collector")
                if isinstance(collector, AttachmentCollector):
                    collector.add_many(result.attachments)
            return result.model_dump()
        except asyncio.TimeoutError:
            return {
                "success": False,
                "error": f"Execution timed out after {payload.timeout_seconds} seconds.",
            }

    return Tool(
        name=TOOL_NAME,
        description=(
            "Create files from LLM-provided content and prepare MinIO upload metadata. "
            "Supported formats: Word(docx), text(json/txt/md...), Excel(xlsx), CSV."
        ),
        parameters=FileWriteToolInput.model_json_schema(),
        handler=_handler,
    )


if __name__ == "__main__":

    async def _demo() -> None:
        tool = create_attachment_file_write_tool()
        session_id = uuid4().hex
        request = AgentRequest(query="demo upload", staff_code="demo")
        args = {
            "files": [
                {
                    "file_name": "hello.txt",
                    "file_type": "txt",
                    "content": "Hello from map_core file write tool.",
                }
            ],
            "minio_object_prefix": f"generated/{session_id}",
            "upload_to_minio": True,
        }
        result = await tool.run(args, request, "-")
        print(json.dumps(result, ensure_ascii=False, indent=2))

    asyncio.run(_demo())
